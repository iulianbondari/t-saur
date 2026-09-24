"""Handling of already-compressed containers (DOCX = ZIP+Deflate, PDF = Flate streams).

Two modes:
  * lossless-bytes: DOCX exploded into members; on extraction it is re-packed and VERIFIED bit-exact
    against the hash of the original. If it does not match (e.g. a DOCX from Word with a different zlib), we fall
    back to storing the original bytes (in production: preflate/precomp for bit-exact).
  * canonical (for agents): the text/structure extracted as Markdown; does NOT reproduce the original bytes.
"""
from __future__ import annotations
import io
import zipfile
from dataclasses import dataclass, field


@dataclass
class ZipRecipe:
    members: list[dict] = field(default_factory=list)
    comment: bytes = b""


def explode_zip(data: bytes) -> tuple[ZipRecipe, dict[str, bytes]]:
    recipe = ZipRecipe()
    parts: dict[str, bytes] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        recipe.comment = z.comment
        for info in z.infolist():
            parts[info.filename] = z.read(info)
            recipe.members.append({
                "name": info.filename,
                "method": info.compress_type,
                "date_time": list(info.date_time),
                "external_attr": info.external_attr,
                "flag_bits": info.flag_bits,
                "create_system": info.create_system,
                "comment": info.comment,
                "extra": info.extra,
            })
    return recipe, parts


def rebuild_zip(recipe: ZipRecipe, parts: dict[str, bytes], level: int | None = None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.comment = recipe.comment
        for m in recipe.members:
            zi = zipfile.ZipInfo(m["name"], date_time=tuple(m["date_time"]))
            zi.compress_type = m["method"]
            zi.external_attr = m["external_attr"]
            zi.create_system = m["create_system"]
            zi.comment = m["comment"]
            zi.extra = m["extra"]
            if level is not None and m["method"] == zipfile.ZIP_DEFLATED:
                zi.compress_level = level
            z.writestr(zi, parts[m["name"]])
    return buf.getvalue()


def try_bitexact_zip(data: bytes):
    """Tries the deflate levels until it gets identical bytes. Returns (recipe, parts, level) or None."""
    try:
        recipe, parts = explode_zip(data)
    except zipfile.BadZipFile:
        return None
    for lvl in (6, 9, 8, 7, 5, 4, 3, 2, 1, None):
        if rebuild_zip(recipe, parts, lvl) == data:
            return recipe, parts, (lvl if lvl is not None else -1)
    return None


def pdf_expand_streams(data: bytes) -> bytes:
    """PDF with all streams decompressed (partial precomp equivalent). NOT bit-exact reversible."""
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    out = doc.tobytes(expand=255, garbage=0, deflate=False)
    doc.close()
    return out


def docx_to_markdown(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for p in d.paragraphs:
        st = (p.style.name or "").lower()
        if st.startswith("heading") or st == "title":
            last = st.split()[-1]
            lvl = 1 if st == "title" else (int(last) if last.isdigit() else 2)
            lines.append("#" * lvl + " " + p.text)
        else:
            lines.append(p.text)
        lines.append("")
    for t in d.tables:
        for r_i, row in enumerate(t.rows):
            cells = [c.text.replace("|", "\\|") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if r_i == 0:
                lines.append("|" + "---|" * len(cells))
        lines.append("")
    return "\n".join(lines)


def pdf_to_markdown(data: bytes) -> str:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    parts: list[str] = []
    for i, page in enumerate(doc):
        parts.append(f"\n<!-- page {i + 1} -->\n")
        parts.append(page.get_text("text"))
    doc.close()
    return "".join(parts)
